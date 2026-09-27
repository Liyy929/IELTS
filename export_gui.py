import tkinter as tk
from tkinter import filedialog, messagebox
import json
import os
from openpyxl import Workbook
from openpyxl.cell.text import InlineFont
from openpyxl.cell.rich_text import TextBlock, CellRichText
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

DATA_FILE = os.path.join(os.path.dirname(__file__), "vocab_data.json")

def load_words():
    if not os.path.exists(DATA_FILE):
        return []
    with open(DATA_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)
    return data.get("words", [])

def make_rich_example(sentence, word):
    """生成富文本例句，将目标单词高亮为红色加粗，并应用 Noto Sans SC Medium 字体"""
    if not sentence or not word:
        return sentence

    lower_sent = sentence.lower()
    lower_word = word.lower()
    idx = lower_sent.find(lower_word)
    if idx == -1:
        return sentence

    # 定义富文本内部的字体
    base_font = InlineFont(rFont="Noto Sans SC Medium", sz=11)
    red_bold = InlineFont(rFont="Noto Sans SC Medium", color="00FF0000", b=True, sz=11)

    blocks = []
    if idx > 0:
        blocks.append(TextBlock(base_font, sentence[:idx]))
    blocks.append(TextBlock(red_bold, sentence[idx:idx + len(word)]))
    if idx + len(word) < len(sentence):
        blocks.append(TextBlock(base_font, sentence[idx + len(word):]))
    
    return CellRichText(blocks)

def export_excel(words_per_column, blank_cols, filename):
    words = load_words()
    if not words:
        messagebox.showwarning("无数据", "生词本为空，请先在浏览器中标记单词。")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "雅思阅读生词"

    # 🎨 定义你要求的字体
    word_font = Font(name='Noto Serif SC Black', size=11)
    example_font = Font(name='Noto Sans SC Medium', size=11)

    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    wrap_align = Alignment(wrap_text=True, vertical="top")

    current_col = 1
    for i in range(0, len(words), words_per_column):
        batch = words[i:i + words_per_column]

        # 写入表头
        headers = ["单词", "词性", "中文释义", "例句"]
        for j, h in enumerate(headers):
            cell = ws.cell(row=1, column=current_col + j, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

        # 写入数据
        for row_idx, w in enumerate(batch, start=2):
            # 1. 单词列（应用 Noto Serif SC Black）
            c1 = ws.cell(row=row_idx, column=current_col, value=w["word"])
            c1.font = word_font  # <--- 这里应用单词字体
            c1.border = thin_border

            # 2. 词性列
            c2 = ws.cell(row=row_idx, column=current_col + 1, value=w.get("pos", ""))
            c2.border = thin_border

            # 3. 中文释义列
            c3 = ws.cell(row=row_idx, column=current_col + 2, value=w.get("meaning", ""))
            c3.border = thin_border
            c3.alignment = wrap_align

            # 4. 例句列（应用 Noto Sans SC Medium）
            example_text = w.get("example", "")
            if example_text:
                rich = make_rich_example(example_text, w["word"])
                c4 = ws.cell(row=row_idx, column=current_col + 3)
                c4.value = rich
                c4.font = example_font  # <--- 这里应用例句字体作为全局回退
                c4.border = thin_border
                c4.alignment = wrap_align
            else:
                c4 = ws.cell(row=row_idx, column=current_col + 3, value="")
                c4.font = example_font
                c4.border = thin_border

        # 设置列宽
        ws.column_dimensions[chr(64 + current_col)].width = 16
        ws.column_dimensions[chr(64 + current_col + 1)].width = 8
        ws.column_dimensions[chr(64 + current_col + 2)].width = 20
        ws.column_dimensions[chr(64 + current_col + 3)].width = 50

        current_col += 4 + blank_cols

    wb.save(filename)
    messagebox.showinfo("导出成功", f"已导出 {len(words)} 个单词到\n{filename}")

def launch_gui():
    root = tk.Tk()
    root.title("雅思生词导出工具")
    root.geometry("420x320")
    root.resizable(False, False)

    tk.Label(root, text="雅思阅读生词导出", font=("Microsoft YaHei", 14, "bold")).pack(pady=10)
    frame = tk.Frame(root)
    frame.pack(pady=15)

    tk.Label(frame, text="每组单词数（一列）:").grid(row=0, column=0, sticky="e", padx=5, pady=5)
    entry_per_col = tk.Entry(frame, width=8)
    entry_per_col.insert(0, "10")
    entry_per_col.grid(row=0, column=1, sticky="w", pady=5)

    tk.Label(frame, text="组间空白列数:").grid(row=1, column=0, sticky="e", padx=5, pady=5)
    entry_blank = tk.Entry(frame, width=8)
    entry_blank.insert(0, "3")
    entry_blank.grid(row=1, column=1, sticky="w", pady=5)

    def on_export():
        try:
            per_col = int(entry_per_col.get())
            blank = int(entry_blank.get())
            if per_col < 1 or blank < 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("输入错误", "请输入有效的正整数。")
            return

        path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel 文件", "*.xlsx")], initialfile="雅思阅读生词.xlsx")
        if path:
            export_excel(per_col, blank, path)

    tk.Button(root, text="📥 一键导出 Excel", font=("Microsoft YaHei", 11), bg="#4472C4", fg="white", width=20, height=2, command=on_export).pack(pady=20)
    
    words = load_words()
    tk.Label(root, text=f"当前生词数: {len(words)}", fg="gray").pack()
    root.mainloop()

if __name__ == "__main__":
    launch_gui()